# -*- coding: utf-8 -*-
"""
core/exporter.py - Phân tích điểm yếu (Diagnostic Analytics)
và Xuất đề thi chuẩn sang file Word (.docx) chuyên nghiệp phục vụ in ấn (Giai đoạn 3).
"""

import os
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from core.storage import StorageManager, get_base_dir


def get_exports_dir() -> Path:
    """Thư mục lưu trữ các tệp đề thi Word xuất ra."""
    exp = get_base_dir() / "exports"
    exp.mkdir(exist_ok=True)
    return exp


def generate_weakness_report(storage: StorageManager) -> Dict[str, Any]:
    """
    Phân tích tỷ lệ đúng/sai theo từng mảng kiến thức,
    xác định điểm yếu cần tập trung cải thiện.
    """
    stats = storage.progress.get("stats", {})
    wrong_map = storage.progress.get("wrong_questions", {})
    
    bank_questions = storage.bank_data.get("questions", [])
    q_type_map = {q["id"]: q.get("type", "single_choice") for q in bank_questions}
    counts = storage.bank_data.get("counts", {})
    
    section_stats = {
        "single_choice": {"name": "Trắc nghiệm 1 đáp án (Phần I)", "total_bank": counts.get("single_choice", 92), "wrong_count": 0},
        "multi_choice": {"name": "Trắc nghiệm nhiều đáp án (Phần II)", "total_bank": counts.get("multi_choice", 57), "wrong_count": 0},
        "true_false": {"name": "Phán đoán Đúng/Sai (Phần III)", "total_bank": counts.get("true_false", 32), "wrong_count": 0},
        "short_answer": {"name": "Trả lời ngắn gọn (Phần IV)", "total_bank": counts.get("short_answer", 12), "wrong_count": 0},
        "case_study": {"name": "Phân tích tình huống (Phần V)", "total_bank": counts.get("case_study", 8), "wrong_count": 0}
    }
    
    for q_id_str, info in wrong_map.items():
        try:
            q_id = int(q_id_str)
        except Exception:
            q_id = q_id_str
        q_type = q_type_map.get(q_id, "single_choice")
        if q_type in section_stats:
            section_stats[q_type]["wrong_count"] += 1

    for k, v in section_stats.items():
        w = v["wrong_count"]
        tot = v["total_bank"]
        mastery = max(0.0, round((tot - w) / tot * 100, 1))
        v["mastery_rate"] = mastery
        if mastery >= 85:
            v["level"] = "VỮNG VÀNG"
            v["status_color"] = "green"
        elif mastery >= 65:
            v["level"] = "TRUNG BÌNH"
            v["status_color"] = "yellow"
        else:
            v["level"] = "CẦN CẢI THIỆN GẤP"
            v["status_color"] = "red"

    sorted_sections = sorted(section_stats.items(), key=lambda x: x[1]["mastery_rate"])
    weakest_section_key = sorted_sections[0][0]
    weakest_section_name = sorted_sections[0][1]["name"]

    top_wrong = sorted(
        wrong_map.values(),
        key=lambda x: x.get("fail_count", 1),
        reverse=True
    )[:5]

    return {
        "total_exams": stats.get("exams_taken", 0),
        "total_answered": stats.get("total_answered", 0),
        "total_correct": stats.get("total_correct", 0),
        "overall_accuracy": stats.get("accuracy_rate", 0.0),
        "total_wrong_active": len(wrong_map),
        "section_breakdown": section_stats,
        "weakest_section_key": weakest_section_key,
        "weakest_section_name": weakest_section_name,
        "top_wrong_questions": top_wrong
    }


def set_cell_background(cell, fill_hex: str):
    """Đổ màu nền cho ô trong bảng docx."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def export_exam_to_docx(
    questions: List[Dict[str, Any]],
    exam_title: str = "Đề Thi Sát Hạch Tổ Trưởng CESBG 2026",
    include_answers: bool = True
) -> Path:
    """
    Xuất danh sách 60 câu hỏi ra tệp Word (.docx) chuyên nghiệp chuẩn format thi cử.
    """
    doc = docx.Document()

    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_top1 = p_top.add_run("TẬP ĐOÀN KHOA HỌC KỸ THUẬT HỒNG HẢI - FOXCONN\n")
    r_top1.bold = True
    r_top1.font.size = Pt(11)
    r_top1.font.color.rgb = RGBColor(70, 70, 70)

    r_top2 = p_top.add_run("CESBG VIỆT NAM - KỲ THI ĐÁNH GIÁ THĂNG CHỨC TỔ TRƯỞNG NĂM 2026\n")
    r_top2.bold = True
    r_top2.font.size = Pt(13)
    r_top2.font.color.rgb = RGBColor(0, 51, 102)

    r_top3 = p_top.add_run(f"BÀI THI SÁT HẠCH LÝ THUYẾT & TÌNH HUỐNG\n({exam_title})\n")
    r_top3.bold = True
    r_top3.font.size = Pt(12)

    r_top4 = p_top.add_run("Thời gian làm bài: 60 phút  |  Tổng điểm: 100 điểm  |  Hình thức: Trắc nghiệm & Tình huống")
    r_top4.italic = True
    r_top4.font.size = Pt(10)

    doc.add_paragraph()

    info_table = doc.add_table(rows=2, cols=4)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    r0 = info_table.rows[0].cells
    r0[0].paragraphs[0].add_run("Họ và tên:").bold = True
    r0[1].paragraphs[0].add_run("................................................")
    r0[2].paragraphs[0].add_run("Mã số thẻ:").bold = True
    r0[3].paragraphs[0].add_run("........................")

    r1 = info_table.rows[1].cells
    r1[0].paragraphs[0].add_run("Xưởng / Bộ phận:").bold = True
    r1[1].paragraphs[0].add_run("................................................")
    r1[2].paragraphs[0].add_run("Điểm số:").bold = True
    r1[3].paragraphs[0].add_run("........... / 100 điểm")

    for row in info_table.rows:
        for cell in row.cells:
            cell.paragraphs[0].runs[0].font.size = Pt(9.5)
            set_cell_background(cell, "F1F5F9")

    doc.add_paragraph()

    p_sec = doc.add_paragraph()
    r_sec = p_sec.add_run("NỘI DUNG ĐỀ THI")
    r_sec.bold = True
    r_sec.font.size = Pt(11)
    r_sec.font.color.rgb = RGBColor(0, 51, 102)

    for idx, q in enumerate(questions, start=1):
        q_type = q.get("type", "single_choice")
        points = q.get("points", 1.0)
        q_text = q.get("question", "")

        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(2)

        r_num = p_q.add_run(f"Câu {idx}: ")
        r_num.bold = True
        r_num.font.size = Pt(10)
        r_num.font.color.rgb = RGBColor(0, 70, 140)

        r_title = p_q.add_run(f"{q_text} ")
        r_title.font.size = Pt(10)

        r_pts = p_q.add_run(f"({points} điểm)")
        r_pts.italic = True
        r_pts.font.size = Pt(9)
        r_pts.font.color.rgb = RGBColor(100, 100, 100)

        options = q.get("options", [])
        if q_type in ("single_choice", "multi_choice") and options:
            for opt in options:
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.left_indent = Inches(0.3)
                p_opt.paragraph_format.space_after = Pt(1)
                r_opt = p_opt.add_run(f"[   ]  {opt}")
                r_opt.font.size = Pt(9.5)

        elif q_type == "true_false":
            p_tf = doc.add_paragraph()
            p_tf.paragraph_format.left_indent = Inches(0.3)
            p_tf.paragraph_format.space_after = Pt(2)
            r_tf = p_tf.add_run("[   ]  Đúng (V)              [   ]  Sai (X)")
            r_tf.bold = True
            r_tf.font.size = Pt(9.5)

        elif q_type == "short_answer":
            p_line = doc.add_paragraph()
            p_line.paragraph_format.left_indent = Inches(0.3)
            p_line.paragraph_format.space_after = Pt(3)
            p_line.add_run("Trả lời: ..........................................................................................................................................\n..........................................................................................................................................................").italic = True

        elif q_type == "case_study":
            p_cs = doc.add_paragraph()
            p_cs.paragraph_format.left_indent = Inches(0.3)
            p_cs.paragraph_format.space_after = Pt(4)
            p_cs.add_run("Phương án xử lý tình huống của thí sinh:\n..........................................................................................................................................................\n..........................................................................................................................................................\n..........................................................................................................................................................").italic = True

    if include_answers:
        doc.add_page_break()

        p_ans_head = doc.add_paragraph()
        p_ans_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_ah = p_ans_head.add_run("PHỤ LỤC: BẢNG ĐÁP ÁN & THANG ĐIỂM CHI TIẾT (ANSWER KEY)")
        r_ah.bold = True
        r_ah.font.size = Pt(13)
        r_ah.font.color.rgb = RGBColor(180, 0, 0)

        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.add_run("(Dành cho Giám thị / Ban Giám khảo chấm thi)").italic = True

        ans_table = doc.add_table(rows=1, cols=4)
        ans_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        hdr_cells = ans_table.rows[0].cells
        hdr_titles = ["Câu số", "Loại câu", "Đáp án chuẩn", "Điểm"]
        for i, title in enumerate(hdr_titles):
            hdr_cells[i].paragraphs[0].add_run(title).bold = True
            hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(9.5)
            set_cell_background(hdr_cells[i], "E2E8F0")

        type_labels = {
            "single_choice": "Trắc nghiệm đơn",
            "multi_choice": "Nhiều đáp án",
            "true_false": "Đúng / Sai",
            "short_answer": "Tự luận ngắn",
            "case_study": "Tình huống"
        }

        for idx, q in enumerate(questions, start=1):
            row_cells = ans_table.add_row().cells
            row_cells[0].paragraphs[0].add_run(str(idx)).bold = True
            row_cells[1].paragraphs[0].add_run(type_labels.get(q.get("type", ""), "Câu hỏi"))
            
            raw_ans = str(q.get("correct_answer", q.get("answer", ""))).strip()
            row_cells[2].paragraphs[0].add_run(raw_ans).bold = True
            row_cells[3].paragraphs[0].add_run(f"{q.get('points', 1.0)}đ")

            for c in row_cells:
                c.paragraphs[0].runs[0].font.size = Pt(9)
            
            if idx % 2 == 0:
                for c in row_cells:
                    set_cell_background(c, "F8FAFC")

        written_qs = [
            (i, q) for i, q in enumerate(questions, start=1)
            if q.get("type") in ("short_answer", "case_study")
        ]
        if written_qs:
            doc.add_paragraph()
            p_det = doc.add_paragraph()
            r_det = p_det.add_run("HƯỚNG DẪN CHẤM ĐIỂM CHI TIẾT PHẦN TỰ LUẬN & TÌNH HUỐNG:")
            r_det.bold = True
            r_det.font.size = Pt(11)

            for idx, q in written_qs:
                p_w = doc.add_paragraph()
                p_w.paragraph_format.left_indent = Inches(0.2)
                p_w.paragraph_format.space_before = Pt(4)
                p_w.add_run(f"Câu {idx}: ").bold = True
                p_w.add_run(f"{q.get('question', '')}\n")
                r_exp = p_w.add_run(f"👉 Đáp án chuẩn: {q.get('correct_answer', q.get('answer', ''))}")
                r_exp.bold = True
                r_exp.font.color.rgb = RGBColor(0, 100, 0)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"De_Thi_CESBG_2026_{timestamp}.docx"
    out_file = get_exports_dir() / filename
    doc.save(str(out_file))
    return out_file
